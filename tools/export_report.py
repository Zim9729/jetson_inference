import pandas as pd
from docx import Document
from docx.shared import Pt, Cm
from docx.oxml.ns import qn
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os
from openpyxl import load_workbook
import re
from PIL import Image
import shutil

# ========== 配置区（已改好） ==========
EXCEL_PATH = r"下行.xlsx"
DATA_SHEET = "数据"
IMG_FOLDER = r"E:\2026081007-1"
IMG_WIDTH = Cm(2.6)
MAX_PX_WIDTH = 700
MAX_PX_HEIGHT = 700

# 故障代码-中文缺陷名称映射
fault_code_map = {
    2: "轨面擦伤",
    3: "轨面掉块",
    16: "扣件缺失",
    24: "弹条断裂",
    32: "弹条移位",
    192: "轨枕开裂\\掉块",
    128: "轨枕开裂\\掉块",
    1024: "道床开裂",
    1536: "道床异物",
    51: "翻浆冒泥",
    48: "扣件地脚螺栓缺失",
    49: "扣件地脚螺栓缺失",
    50: "扣件地脚螺栓垫片松动",
    8192: "扣件地脚螺栓垫片松动"
}
# ========================================

def set_font(run, font_name="宋体", size=12, bold=False):
    run.font.name = font_name
    run.font.size = Pt(size)
    run.font.bold = bold
    r = run._element
    r.rPr.rFonts.set(qn('w:eastAsia'), font_name)

# 单元格 水平+垂直 双向居中
def set_cell_full_center(cell):
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    for paragraph in cell.paragraphs:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

# 毫米值自动转换标准K里程
# -30324750 → K-30-324.75
# 4071750 → K4+071.75
def mm_to_k_mile(mm_val):
    try:
        mm = int(float(mm_val))
    except:
        return str(mm_val)
    neg = mm < 0
    mm_abs = abs(mm)
    km = mm_abs // 1000000
    rem = mm_abs % 1000000
    m = rem / 1000.0
    m_str = f"{m:06.2f}"
    if neg:
        return f"K-{km}-{m_str}"
    else:
        return f"K{km}+{m_str}"

# 图片压缩+异常捕获，损坏图片直接跳过不卡死
def compress_image(origin_path, temp_path, max_w, max_h):
    try:
        im = Image.open(origin_path)
        im.thumbnail((max_w, max_h))
        im.save(temp_path, "JPEG", quality=65, optimize=True)
        return True
    except Exception as e:
        print(f"图片损坏跳过：{origin_path} | {e}")
        return False

# 读取Excel
if not os.path.exists(EXCEL_PATH):
    print(f"错误：文件 {EXCEL_PATH} 不存在")
    exit()
if not os.path.isdir(IMG_FOLDER):
    print(f"警告：图片文件夹 {IMG_FOLDER} 不存在，图片全部显示缺失")

wb = load_workbook(EXCEL_PATH, data_only=True)
ws_data = wb[DATA_SHEET]
all_data = []

for row in range(2, ws_data.max_row + 1):
    loc_mm = ws_data.cell(row=row, column=35).value
    code = ws_data.cell(row=row, column=37).value
    area_raw = ws_data.cell(row=row, column=40).value
    img_text = ws_data.cell(row=row, column=41).value

    # 区间统一替换 到 → -
    if pd.isna(area_raw) or str(area_raw).strip() == "":
        area = "未知区间"
    else:
        area = str(area_raw).strip().replace("到", "-")

    # 缺陷名称转换
    try:
        code_num = int(code)
        defect_name = fault_code_map.get(code_num, f"未知故障{code_num}")
    except:
        defect_name = f"无效代码{code}"

    # 图片路径解析
    img_path = ""
    if img_text is not None:
        raw = str(img_text).replace("\n","").replace(" ","").strip()
        fn = os.path.basename(raw)
        if fn.endswith(".jpg"):
            img_path = os.path.join(IMG_FOLDER, fn)

    mile_text = mm_to_k_mile(loc_mm)
    row_data = {
        "缺陷名称": defect_name,
        "里程位置": mile_text,
        "区间": area,
        "图片完整路径": img_path
    }
    all_data.append(row_data)
wb.close()

df = pd.DataFrame(all_data)
print("总数据行数：", len(df))
if len(df) == 0:
    print("无故障数据，程序退出")
    exit()
group_defect = df.groupby("缺陷名称")

# 临时图片缓存文件夹
temp_dir = "temp_img_cache"
if not os.path.exists(temp_dir):
    os.mkdir(temp_dir)

# 每个缺陷单独生成Word
for defect_name, group_df in group_defect:
    safe_name = defect_name.replace("\\", "、")
    word_file = f"{safe_name}.docx"
    print(f"\n===== 生成文档 {word_file} 共{len(group_df)}条 =====")
    doc = Document()

    # 文档大标题
    title = doc.add_heading(f"{defect_name}缺陷统计明细报告", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(title.runs[0], "黑体", 18, True)

    # ========== 已删除首页多余说明段落 ==========

    # 一、区间汇总目录（完全沿用Excel原始出现顺序，不自动重排）
    doc.add_paragraph()
    h_summary = doc.add_heading("一、区间记录汇总目录", level=1)
    set_font(h_summary.runs[0], "黑体", 14, True)
    area_count = group_df["区间"].value_counts().to_dict()
    area_order_list = group_df["区间"].drop_duplicates().tolist()
    for area in area_order_list:
        cnt = area_count[area]
        p_dir = doc.add_paragraph(f"▶ {area}（{cnt}条记录）")
        set_font(p_dir.runs[0])
    doc.add_page_break()

    # 区间分组，sort=False 保持Excel原生顺序，不自动字母排序
    area_group = group_df.groupby("区间", sort=False)
    for area_name, area_data in area_group:
        sub_title = doc.add_heading(f"{area_name}（{len(area_data)}条记录）", level=2)
        set_font(sub_title.runs[0], "黑体", 13, True)
        # 创建4列表格
        table = doc.add_table(rows=1, cols=4)
        table.style = "Table Grid"
        table.alignment = WD_ALIGN_PARAGRAPH.CENTER
        headers = ["里程位置", "缺陷名称", "所属区间", "故障实拍图"]
        h_cells = table.rows[0].cells
        for idx, cell in enumerate(h_cells):
            cell.text = headers[idx]
            set_cell_full_center(cell)
            set_font(cell.paragraphs[0].runs[0], bold=True)
        # 区间内部完全沿用Excel行顺序，不做里程二次排序（你Excel已排好）
        for _, row in area_data.iterrows():
            r_cells = table.add_row().cells
            r_cells[0].text = row["里程位置"]
            r_cells[1].text = row["缺陷名称"]
            r_cells[2].text = row["区间"]
            img_cell = r_cells[3]
            img_p = row["图片完整路径"]
            print("处理图片：", img_p)
            if img_p and os.path.exists(img_p):
                fn = os.path.basename(img_p)
                temp_p = os.path.join(temp_dir, fn)
                if compress_image(img_p, temp_p, MAX_PX_WIDTH, MAX_PX_HEIGHT):
                    para = img_cell.paragraphs[0]
                    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    para.add_run().add_picture(temp_p, width=IMG_WIDTH)
                    p2 = img_cell.add_paragraph()
                    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    set_font(p2.add_run(f"图:{fn}"), size=9)
                    os.remove(temp_p)
                else:
                    img_cell.text = "图片损坏，无法加载"
            else:
                img_cell.text = f"图片缺失\n{img_p}"
            for c in r_cells:
                set_cell_full_center(c)
        doc.add_page_break()
    doc.save(word_file)
    print(f"✅ {word_file} 生成完成")

# 清理临时缓存
if os.path.exists(temp_dir):
    shutil.rmtree(temp_dir)
print("\n===== 全部文档生成完毕 =====")
print("功能汇总：1.毫米自动转标准K里程 2.表格全部单元格双向居中 3.区间/明细完全沿用Excel你排好的顺序不乱 4.无首页多余文字 5.图片自动加载、损坏不卡死 6.分缺陷单独Word")